"""Извлечение текста из PDF, DOCX, XLSX. Запускается отдельным процессом (python -m wshub.extract),
чтобы зависший разбор можно было убить по таймауту."""
from __future__ import annotations

import csv
import io
import json
import sys
from pathlib import Path

EXTRACTABLE = (".pdf", ".docx", ".xlsx", ".xlsm")


class Out:
    """Накопитель текста с лимитом в байтах."""

    def __init__(self, budget: int):
        self.budget, self.used, self.parts, self.full = budget, 0, [], False

    def add(self, s: str) -> bool:
        if self.full:
            return False
        b = len(s.encode("utf-8")) + 1
        if self.used + b > self.budget:
            room = self.budget - self.used
            self.parts.append(s.encode("utf-8")[: max(room, 0)].decode("utf-8", "ignore"))
            self.full = True
            return False
        self.parts.append(s)
        self.used += b
        return True

    def text(self) -> str:
        return "\n".join(self.parts)


def _pdf(path: Path, out: Out) -> str:
    from pypdf import PdfReader

    r = PdfReader(path)
    if r.is_encrypted and not r.decrypt(""):
        raise ValueError("PDF зашифрован паролем")
    total = len(r.pages)
    for i, page in enumerate(r.pages, 1):
        if not (out.add(f"--- страница {i} из {total} ---") and out.add(page.extract_text() or "")):
            return f"на странице {i} из {total}"
    return ""


def _docx(path: Path, out: Out) -> str:
    import docx
    from docx.table import Table

    d = docx.Document(str(path))
    for n, block in enumerate(d.iter_inner_content(), 1):
        if isinstance(block, Table):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in block.rows]
            ok = out.add("[таблица]") and all(out.add(r) for r in rows) and out.add("[/таблица]")
        else:
            ok = out.add(block.text)
        if not ok:
            return f"на блоке {n} документа"
    return ""


def _xlsx(path: Path, out: Out) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if not out.add(f"--- лист {ws.title} ---"):
                return f"на листе {ws.title}"
            empty = 0  # пустые строки выводим, только если за ними есть данные
            for n, row in enumerate(ws.iter_rows(values_only=True), 1):
                cells = ["" if v is None else str(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells:
                    empty += 1
                    continue
                buf = io.StringIO()
                csv.writer(buf, lineterminator="").writerow(cells)
                if not (all(out.add("") for _ in range(empty)) and out.add(buf.getvalue())):
                    return f"на листе {ws.title}, строка {n}"
                empty = 0
    finally:
        wb.close()
    return ""


def extract(path: Path, budget: int) -> dict:
    ext = path.suffix.lower()
    out = Out(budget)
    where = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx}[ext](path, out)
    return {"text": out.text(), "truncated": where}


def main() -> None:
    path, budget = Path(sys.argv[1]), int(sys.argv[2])
    try:
        res = extract(path, budget)
    except Exception as e:  # отдаём текст ошибки родителю, без трассировки
        res = {"error": f"{type(e).__name__}: {e}"}
    sys.stdout.write(json.dumps(res, ensure_ascii=False))


if __name__ == "__main__":
    main()
