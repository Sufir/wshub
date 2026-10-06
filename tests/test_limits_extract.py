"""read/extract: заголовок, бинарные файлы, обрезка по лимитам; extract для PDF, DOCX, XLSX."""
import pytest

from wshub import core
from wshub.core import WsError


def make_pdf(path, text: str) -> None:
    """Минимальный PDF с одной страницей и текстом Helvetica."""
    content = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(out)


def make_docx(path) -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Первый абзац отчёта")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ключ", "значение"
    t.cell(1, 0).text, t.cell(1, 1).text = "порт", "8080"
    d.add_paragraph("Последний абзац")
    d.save(path)


def make_xlsx(path, rows: int = 3) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Хосты"
    ws.append(["host", "ip", "note"])
    for i in range(rows - 1):
        ws.append([f"srv{i}", f"10.0.0.{i}", "a, b"])
    s2 = wb.create_sheet("Итог")
    s2.append(["всего", rows - 1])
    s2.append([])
    s2.append(["конец"])
    wb.save(path)


def test_read_header_and_numbers(env):
    ws = env.open()
    out = env.hub.read(ws, "a.txt")
    assert out.splitlines()[0] == "=== содержимое файла a.txt ==="
    assert "1\thello\n2\tworld" in out
    assert out.endswith("=== конец файла ===")
    out = env.hub.read(ws, "a.txt", offset=2)
    assert "1\thello" not in out and "2\tworld" in out


def test_read_line_limit(env):
    (env.proj / "many.txt").write_text("".join(f"line{i}\n" for i in range(1, 101)))
    ws = env.open()
    out = env.hub.read(ws, "many.txt", offset=10, limit=5)
    assert "10\tline10" in out and "14\tline14" in out and "15\tline15" not in out
    assert "offset=15" in out


def test_read_size_limit(env):
    env.write_registry(max_read_kb=1)
    (env.proj / "big.txt").write_text("".join(f"{'x' * 90} {i}\n" for i in range(1, 200)))
    ws = env.open()
    out = env.hub.read(ws, "big.txt")
    assert "обрезано по лимиту 1 КБ" in out and "offset=" in out
    assert len(out.encode()) < 1024 + 300
    (env.proj / "oneline.min.js").write_text("y" * 5000)
    out = env.hub.read(ws, "oneline.min.js")
    assert "[строка обрезана]" in out and "offset=2" in out


def test_read_binary_refused(env):
    (env.proj / "img.png").write_bytes(b"\x89PNG\r\n\x1a\n\x00\x00\x00")
    make_pdf(env.proj / "doc.pdf", "Hi")
    ws = env.open()
    with pytest.raises(WsError, match="бинарный.*extract"):
        env.hub.read(ws, "img.png")
    with pytest.raises(WsError, match="используй extract"):
        env.hub.read(ws, "doc.pdf")


def test_extract_pdf(env):
    make_pdf(env.proj / "Report.PDF", "Hello wshub PDF")
    ws = env.open()
    out = env.hub.extract(ws, "Report.PDF")
    assert out.startswith("=== содержимое файла Report.PDF ===")
    assert "--- страница 1 из 1 ---" in out and "Hello wshub PDF" in out


def test_extract_docx(env):
    make_docx(env.proj / "r.docx")
    out = env.hub.extract(env.open(), "r.docx")
    assert out.startswith("=== содержимое файла r.docx ===")
    i1, i2, i3 = out.index("Первый абзац"), out.index("порт | 8080"), out.index("Последний абзац")
    assert i1 < i2 < i3
    assert "ключ | значение" in out


def test_extract_xlsx(env):
    make_xlsx(env.proj / "hosts.xlsx")
    out = env.hub.extract(env.open(), "hosts.xlsx")
    assert out.startswith("=== содержимое файла hosts.xlsx ===")
    assert "--- лист Хосты ---\nhost,ip,note\nsrv0,10.0.0.0,\"a, b\"\nsrv1,10.0.0.1,\"a, b\"" in out
    assert "--- лист Итог ---\nвсего,2\n\nконец" in out


def test_extract_truncated(env):
    env.write_registry(max_read_kb=1)
    make_xlsx(env.proj / "big.xlsx", rows=500)
    out = env.hub.extract(env.open(), "big.xlsx")
    assert "обрезано по лимиту 1 КБ на листе Хосты" in out
    assert "srv400" not in out


def test_extract_errors(env):
    ws = env.open()
    with pytest.raises(WsError, match="read"):
        env.hub.extract(ws, "a.txt")
    (env.proj / "broken.pdf").write_bytes(b"not a pdf at all")
    with pytest.raises(WsError, match="не удалось извлечь"):
        env.hub.extract(ws, "broken.pdf")


def test_extract_timeout(env, monkeypatch):
    make_pdf(env.proj / "d.pdf", "x")
    monkeypatch.setattr(core, "EXTRACT_TIMEOUT", 0.001)
    with pytest.raises(WsError, match="не уложился"):
        env.hub.extract(env.open(), "d.pdf")


def test_grep_limit(env, engine):
    (env.proj / "g.txt").write_text("match\n" * 400)
    out = env.hub.grep(env.open(), "match")
    lines = out.splitlines()
    assert len([ln for ln in lines if ln.startswith("g.txt:")]) == 300
    assert "обрезано на 300 совпадениях" in lines[-1]


def test_grep_timeout(env, engine, monkeypatch):
    for i in range(50):
        (env.proj / f"f{i}.txt").write_text("needle\n" * 50)
    monkeypatch.setattr(core, "GREP_TIMEOUT", 0)
    out = env.hub.grep(env.open(), "needle")
    assert "таймауту" in out or "обрезано" in out


def test_grep_bad_regex(env):
    with pytest.raises(WsError, match="регулярное"):
        env.hub.grep(env.open(), "(")


def test_tree_limits(env):
    d = env.proj / "many"
    d.mkdir()
    for i in range(2100):
        (d / f"f{i:04}.txt").write_text("")
    ws = env.open()
    out = env.hub.tree(ws, ".", 2)
    assert len(out.splitlines()) == 2001
    assert "обрезано на 2000 строках" in out.splitlines()[-1]
    (env.proj / "a1" / "a2" / "a3" / "a4" / "a5" / "a6").mkdir(parents=True)
    out = env.hub.tree(ws, "a1", 9)  # depth 4 от a1: a2…a5
    assert "a5/" in out and "a6/" not in out and "depth уменьшен до 4" in out


def test_tree_hides_git_and_denied_contents(env):
    (env.proj / ".git" / "objects").mkdir(parents=True)
    (env.proj / "node_modules" / "pkg").mkdir(parents=True)
    out = env.hub.tree(env.open(), ".", 4)
    assert ".git/  [не раскрыт]" in out and "node_modules/  [deny]" in out
    assert "objects" not in out and "pkg" not in out
